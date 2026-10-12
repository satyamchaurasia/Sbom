"""Build the Word manual from docs/SBOM_Fixer_Rule_Book.md.

Writes docs/SBOM_Fixer_Rule_Book.docx and the copy SBOM_Fixer_Rule_Book_And_Installation_guide.docx in the
repository root. Needs python-docx (not a dependency of the tool):

    pip install python-docx
    python tools/make_rule_catalog.py      # refresh the rule table first
    python tools/make_rule_book_docx.py

Supports the Markdown the manual uses: # title, ## / ### headings, paragraphs with **bold** and `code`, - and 1.
lists, | tables |, ``` code blocks, > callouts and --- separators.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "SBOM_Fixer_Rule_Book.md"
TARGETS = [ROOT / "docs" / "SBOM_Fixer_Rule_Book.docx", ROOT / "SBOM_Fixer_Rule_Book_And_Installation_guide.docx"]

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CODE_FONT = "Consolas"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def shade(element, fill: str) -> None:
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "A6A6A6")
        el.append(b)
    tbl_pr.append(el)


def add_inline(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
            run.bold = bold
        else:
            run = paragraph.add_run(part)
            run.bold = bold
        if size:
            run.font.size = Pt(size)


def toc(doc) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        elif text:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)
        else:
            t = OxmlElement("w:t")
            t.text = "Right-click here and choose Update Field to build the table of contents."
            run._r.append(t)


def split_row(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|") for c in cells]


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], [r for r in rows[2:]]
    wide = len(header) >= 6
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    size = 7.5 if wide else 9
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        shade(cell._tc, "1F4E79")
        p = cell.paragraphs[0]
        add_inline(p, text, size, bold=True)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r, row in enumerate(body, start=1):
        for i in range(len(header)):
            cell = table.rows[r].cells[i]
            if r % 2 == 0:
                shade(cell._tc, "F2F6FA")
            add_inline(cell.paragraphs[0], row[i] if i < len(row) else "", size)
    if header[0] == "Rule" and len(header) == 7:  # the rule catalog: room for the IDs and the texts
        table.autofit = False
        for i, width in enumerate((1.9, 0.9, 1.5, 1.6, 3.0, 5.5, 2.6)):
            for row_ in table.rows:
                row_.cells[i].width = Cm(width)
    doc.add_paragraph()


def add_code(doc, lines: list[str]) -> None:
    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        shade(p._p, "F2F2F2")
        fmt = p.paragraph_format
        fmt.space_before = Pt(4 if i == 0 else 0)
        fmt.space_after = Pt(4 if i == len(lines) - 1 else 0)
        fmt.left_indent = Cm(0.3)
        run = p.add_run(line or " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(8.5)


def build() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.color.rgb = ACCENT
        st.font.size = Pt(size)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title")
        elif line.startswith("## "):
            title = line[3:].strip()
            if title == "Contents":
                doc.add_heading(title, level=1)
                toc(doc)
                i += 1
                while i < len(lines) and not lines[i].startswith("---"):  # the inline list is replaced by the TOC
                    i += 1
                continue
            if not first_h1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h1 = False
            doc.add_heading(title, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif stripped == "---":
            pass
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            shade(p._p, "FFF4CE")
            p.paragraph_format.left_indent = Cm(0.3)
            add_inline(p, stripped[2:])
        elif re.match(r"^\s*[-*] ", line):
            level = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            # written numbers instead of Word's List Number, which keeps counting across the whole document
            number, text = re.match(r"^\s*(\d+)\. (.*)$", line).groups()  # type: ignore[union-attr]
            while i + 1 < len(lines) and lines[i + 1].startswith("   ") and lines[i + 1].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.) ", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{number}.\t{text}")
        elif stripped:
            text = stripped
            while i + 1 < len(lines) and lines[i + 1].startswith("   ") and lines[i + 1].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.) ", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            add_inline(doc.add_paragraph(), text)
        i += 1
    return doc


def main() -> None:
    doc = build()
    doc.save(TARGETS[0])
    for extra in TARGETS[1:]:
        shutil.copyfile(TARGETS[0], extra)
    print("written: " + ", ".join(str(t.relative_to(ROOT)) for t in TARGETS))


if __name__ == "__main__":
    main()
